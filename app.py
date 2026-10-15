import io
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import streamlit as st


@st.cache_data
def consolidar_archivos(file_nequi, file_movs):
  # 1. Leer documento.xlsx (Nequi) detectando dinámicamente la fila de cabecera
  df1_raw = pd.read_excel(file_nequi, sheet_name='Table 1')
  header_idx = 7
  for idx, row in df1_raw.iterrows():
    if row.astype(str).str.contains('Fecha del movimiento').any():
      header_idx = idx
      break

  df1 = pd.read_excel(
      file_nequi, sheet_name='Table 1', skiprows=header_idx + 1
  )
  df1 = df1.iloc[:, [0, 1, 3]]
  df1.columns = ['Fecha', 'Descripción', 'Valor']

  # 2. Leer el segundo archivo de movimientos (.xls)
  df2 = pd.read_excel(file_movs, sheet_name='Movimientos')
  val_debito = (
      df2['Valor Débito']
      if 'Valor Débito' in df2.columns
      else df2.iloc[:, 2]
  )
  val_credito = (
      df2['Valor Crédito']
      if 'Valor Crédito' in df2.columns
      else df2.iloc[:, 3]
  )
  df2['Valor'] = val_debito.fillna(0) + val_credito.fillna(0)
  df2 = df2[['Fecha', 'Descripción', 'Valor']]

  # 3. Unir ambos DataFrames
  df_total = pd.concat([df1, df2], ignore_index=True)
  df_total['Valor'] = pd.to_numeric(df_total['Valor'], errors='coerce').fillna(
      0
  )

  # 4. Filtrar solo salidas de dinero (valores negativos)
  df_salidas = df_total[df_total['Valor'] < 0].copy()

  # 5. Limpiar descripciones para extraer el nombre de la persona
  limpiar_textos = (
      df_salidas['Descripción']
      .astype(str)
      .str.replace('Para ', '', case=False)
      .str.replace('ENVIO CON BRE-B A: ', '', case=False)
      .str.strip()
  )
  df_salidas['Persona'] = limpiar_textos
  df_salidas['Valor_Absoluto'] = df_salidas['Valor'].abs()

  # 6. Agrupar por Persona: sumar los valores y contar los movimientos
  df_consolidado = (
      df_salidas.groupby('Persona')
      .agg(
          Total_Enviado=('Valor_Absoluto', 'sum'),
          Cantidad_Envios=('Valor_Absoluto', 'count'),
      )
      .reset_index()
  )

  # Ordenar de mayor a menor monto
  df_consolidado = df_consolidado.sort_values(
      by='Total_Enviado', ascending=False
  ).reset_index(drop=True)

  return df_consolidado


def generar_excel_bytes(df):
  """Genera un archivo Excel formateado profesionalmente en memoria y retorna sus bytes"""
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = 'Consolidado Envíos'

  # Asegurar visibilidad de líneas de cuadrícula
  ws.views.sheetView[0].showGridLines = True

  # Título principal
  ws.merge_cells('A1:C1')
  ws['A1'] = 'CONSOLIDADO DE MOVIMIENTOS Y ENVÍOS POR PERSONA'
  ws['A1'].font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')
  ws['A1'].fill = PatternFill(
      start_color='1F4E78', end_color='1F4E78', fill_type='solid'
  )
  ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
  ws.row_dimensions[1].height = 35

  # Subtítulo
  ws.merge_cells('A2:C2')
  ws['A2'] = 'Resumen agrupado sin nombres repetidos'
  ws['A2'].font = Font(name='Calibri', size=10, italic=True, color='595959')
  ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
  ws.row_dimensions[2].height = 20

  # Cabeceras
  headers = ['Persona / Destinatario', 'Total Enviado', 'Cantidad de Envíos']
  ws.row_dimensions[4].height = 24

  header_fill = PatternFill(
      start_color='2F5597', end_color='2F5597', fill_type='solid'
  )
  header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
  header_border = Border(
      left=Side(style='thin', color='BFBFBF'),
      right=Side(style='thin', color='BFBFBF'),
      top=Side(style='thin', color='BFBFBF'),
      bottom=Side(style='medium', color='1F4E78'),
  )

  for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=4, column=col_num, value=header)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(
        horizontal='center' if col_num > 1 else 'left', vertical='center'
    )
    cell.border = header_border

  # Datos
  row_start = 5
  thin_border = Border(
      left=Side(style='thin', color='D9D9D9'),
      right=Side(style='thin', color='D9D9D9'),
      top=Side(style='thin', color='D9D9D9'),
      bottom=Side(style='thin', color='D9D9D9'),
  )
  zebra_fill = PatternFill(
      start_color='F9FAFB', end_color='F9FAFB', fill_type='solid'
  )

  for i, row_data in df.iterrows():
    current_row = row_start + i
    ws.row_dimensions[current_row].height = 20

    cell_persona = ws.cell(row=current_row, column=1, value=row_data['Persona'])
    cell_total = ws.cell(
        row=current_row, column=2, value=row_data['Total_Enviado']
    )
    cell_cantidad = ws.cell(
        row=current_row, column=3, value=row_data['Cantidad_Envios']
    )

    cell_persona.alignment = Alignment(horizontal='left', vertical='center')
    cell_total.alignment = Alignment(horizontal='right', vertical='center')
    cell_cantidad.alignment = Alignment(horizontal='center', vertical='center')

    cell_total.number_format = '$#,##0.00'
    cell_cantidad.number_format = '#,##0'

    for col_idx in range(1, 4):
      c = ws.cell(row=current_row, column=col_idx)
      c.font = Font(name='Calibri', size=10)
      c.border = thin_border
      if i % 2 == 1:
        c.fill = zebra_fill

  # Fila de Totales
  total_row = row_start + len(df)
  ws.row_dimensions[total_row].height = 22

  cell_t_label = ws.cell(row=total_row, column=1, value='TOTAL GENERAL')
  cell_t_sum = ws.cell(
      row=total_row, column=2, value=f'=SUM(B{row_start}:B{total_row-1})'
  )
  cell_t_count = ws.cell(
      row=total_row, column=3, value=f'=SUM(C{row_start}:C{total_row-1})'
  )

  total_fill = PatternFill(
      start_color='D9E1F2', end_color='D9E1F2', fill_type='solid'
  )
  total_font = Font(name='Calibri', size=11, bold=True, color='000000')
  total_border = Border(
      top=Side(style='thin', color='1F4E78'),
      bottom=Side(style='double', color='1F4E78'),
      left=Side(style='thin', color='D9D9D9'),
      right=Side(style='thin', color='D9D9D9'),
  )

  for col_idx in range(1, 4):
    c = ws.cell(row=total_row, column=col_idx)
    c.fill = total_fill
    c.font = total_font
    c.border = total_border

  cell_t_label.alignment = Alignment(horizontal='left', vertical='center')
  cell_t_sum.alignment = Alignment(horizontal='right', vertical='center')
  cell_t_sum.number_format = '$#,##0.00'
  cell_t_count.alignment = Alignment(horizontal='center', vertical='center')
  cell_t_count.number_format = '#,##0'

  # Ajustar ancho de columnas automáticamente
  for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 5, 18)

  # Guardar en buffer de memoria
  output = io.BytesIO()
  wb.save(output)
  output.seek(0)
  return output.getvalue()


# --- Interfaz en Streamlit ---
st.title('Consolidador de Movimientos a Excel')

uploaded_nequi = st.file_uploader('Sube documento.xlsx (Nequi)', type=['xlsx'])
uploaded_movs = st.file_uploader(
    'Sube Movs_CA_15092026_15102026.xls', type=['xls', 'xlsx']
)

if uploaded_nequi and uploaded_movs:
  df_resultado = consolidar_archivos(uploaded_nequi, uploaded_movs)

  st.success('¡Consolidado listo para descargar en Excel!')

  # Mostrar vista previa interactiva en la app
  st.dataframe(df_resultado, use_container_width=True)

  # Generar el archivo Excel descargable
  excel_bytes = generar_excel_bytes(df_resultado)

  st.download_button(
      label='📥 Descargar Excel Consolidado (.xlsx)',
      data=excel_bytes,
      file_name='movimientos_consolidados_por_persona.xlsx',
      mime=(
          'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
      ),
  )
